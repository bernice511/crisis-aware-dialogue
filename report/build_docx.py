"""Generate the 'Datasets and Exploratory Analysis' section as a one-page Word doc.

Produces report/data_analysis.docx: condensed prose, one figure, one comparison
table, sized to fit a single page so a teammate can paste it into the report.

Run:
    python report/build_docx.py
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
FIG = HERE / "figures"
OUT = HERE / "data_analysis.docx"

doc = Document()
sec = doc.sections[0]
for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, side, Inches(0.6))

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.0

GRAY = RGBColor(0x55, 0x55, 0x55)


def title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)


def body(lead, rest):
    """Paragraph with a bold lead-in phrase, e.g. 'CRADLEBench. ...'."""
    p = doc.add_paragraph()
    if lead:
        b = p.add_run(lead + " ")
        b.bold = True
    p.add_run(rest)
    return p


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = GRAY


# ---------------------------------------------------------------------------
title("Datasets and Exploratory Analysis")

body("", "We use two complementary public datasets. CRADLEBench (Byun et al., 2025) "
     "provides real, clinician-annotated crisis labels over help-seeking posts, while "
     "DeepSuiMind (Li et al., 2025) provides synthetic dialogues that isolate implicit "
     "suicidal ideation. Our exploratory analysis examines label structure, class "
     "imbalance, text length, and data quality across both.")

body("CRADLEBench.",
     "8,259 help-seeking posts (train 7,239 / dev 420 / test 600) in a multi-label "
     "task tagging each post with <crisis_type>_<temporal> or no_crisis; it is the "
     "first crisis benchmark to add temporal (ongoing/past) labels. Labels are "
     "strongly imbalanced (Figure 1): no_crisis dominates (2,500 occurrences), "
     "followed by self-harm (1,715), passive (1,171) and active (909) suicide "
     "ideation, domestic violence (921), rape (800), sexual harassment (692), and "
     "child abuse/endangerment (311); most posts carry one label but up to four. The "
     "temporal skew is clinically informative — 4,262 ongoing vs. 2,257 past overall, "
     "with self-harm and suicide ideation mostly ongoing but rape and child abuse "
     "mostly past — making temporal state a first-class modelling target. Self-harm "
     "and suicide ideation co-occur most often. The raw labels required normalizing "
     "245 malformed entries (e.g. 'No crisis', 'suicideideation(passive)'), cutting "
     "the vocabulary from 22 to 15 tags.")

# One key figure, kept small to preserve the one-page budget.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
p.add_run().add_picture(str(FIG / "cradle_crisis_types.png"), width=Inches(3.6))
caption("Figure 1: CRADLEBench crisis-type frequency (all splits, normalized labels).")

body("DeepSuiMind.",
     "Fully synthetic: 1,605 narratives (median 306 words; no missing values or "
     "duplicates), each tagged with a Scenario (11 situations, led by Depression and "
     "Loneliness), a Negative Core Belief (~11 cognitive distortions, led by Emotional "
     "Reasoning), and a D/S-IAT Intention Category. After fixing hyphen/space spelling "
     "variants (268 rows), Intention resolves to three balanced classes: Death-Me "
     "(537), Death-Not-Me (534), and Life-Not-Me (534).")

body("Comparison and implications.",
     "The datasets differ in unit, label scheme, and provenance (Table 1): "
     "CRADLEBench is real, multi-label, and shorter/variable in length; DeepSuiMind is "
     "synthetic, categorical, and length-bounded. The real-vs-synthetic domain gap and "
     "the class imbalance directly shape our preprocessing, imbalanced-data "
     "optimization, and evaluation design.")

# Compact comparison table
t = doc.add_table(rows=1, cols=3)
t.style = "Light Grid Accent 1"
hdr = ["", "CRADLEBench", "DeepSuiMind"]
for j, htext in enumerate(hdr):
    r = t.rows[0].cells[j].paragraphs[0].add_run(htext)
    r.bold = True
    r.font.size = Pt(9)
rows = [["Rows", "8,259", "1,605"],
        ["Unit", "help-seeking post", "synthetic narrative"],
        ["Task", "multi-label crisis + temporal", "categorical (scenario/belief/intent)"],
        ["Median words", "143", "306"],
        ["Provenance", "real, clinician-annotated", "fully synthetic"]]
for row in rows:
    cells = t.add_row().cells
    for j, val in enumerate(row):
        r = cells[j].paragraphs[0].add_run(str(val))
        r.font.size = Pt(9)
for row in t.rows:
    row.cells[0].width = Inches(1.1)
    row.cells[1].width = Inches(2.6)
    row.cells[2].width = Inches(3.0)
caption("Table 1: Summary comparison of the two datasets.")

# Compact references
refp = doc.add_paragraph()
refp.paragraph_format.space_before = Pt(2)
rr = refp.add_run(
    "References: Byun et al. (2025), CRADLE Bench, arXiv:2510.23845.  "
    "Li et al. (2025), Can LLMs Identify Implicit Suicidal Ideation?, "
    "Findings of EMNLP 2025, arXiv:2502.17899.")
rr.font.size = Pt(8)
rr.font.color.rgb = GRAY

doc.save(str(OUT))
print(f"Wrote {OUT}")
